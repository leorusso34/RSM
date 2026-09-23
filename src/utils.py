#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Utility di formattazione, calcolo e I/O per RSM."""

import csv
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from typing import Any, Optional, Tuple, List, Dict

from .config import STATUS_PREORDER, STATUS_SOLD, STATUS_STOCK, PROJECT_DIR, MONTHS_SHORT_IT

def parse_money(value: Any, allow_empty: bool = False) -> Optional[Decimal]:
    """Converte un importo monetario valido in Decimal con due decimali."""
    if value is None or str(value).strip() == "":
        if allow_empty:
            return None
        raise ValueError("Valore monetario obbligatorio.")

    text = str(value).strip().replace("€", "").replace("$", "").replace("£", "")
    text = text.replace(" ", "")
    if not re.fullmatch(r"-?\d+(?:[,.]\d{1,2}|[.]\d{3}(?:[,.]\d{1,2})?)?", text):
        if allow_empty:
            return None
        raise ValueError("Inserisci un importo valido con massimo due decimali.")

    if not text or text in {"-", ".", ","}:
        if allow_empty:
            return None
        raise ValueError("Inserisci un importo valido.")

    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text and len(text.rsplit(",", 1)[1]) == 3:
        text = text.replace(",", "")
    elif "." in text and len(text.rsplit(".", 1)[1]) == 3:
        text = text.replace(".", "")
    else:
        text = text.replace(",", ".")

    try:
        amount = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError("Inserisci un importo valido.") from exc

    if amount < 0:
        raise ValueError("L'importo non può essere negativo.")

    return amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def money_to_cents(value: Optional[Any]) -> Optional[int]:
    """Converte un importo Decimal in centesimi interi per SQLite."""
    if value is None:
        return None
    amount = value if isinstance(value, Decimal) else Decimal(str(value))
    return int((amount * 100).quantize(Decimal("1"), rounding=ROUND_HALF_UP))


def cents_to_money(value: Optional[int]) -> Optional[Decimal]:
    """Converte centesimi interi SQLite in un importo Decimal."""
    if value is None:
        return None
    return (Decimal(value) / Decimal("100")).quantize(Decimal("0.01"))

def money(value: Any) -> str:
    """Formatta un numero come importo monetario in Euro (€ 1.234,56)."""
    if value is None:
        return "€ 0,00"
    try:
        amount = Decimal(str(value)).quantize(Decimal("0.01"))
        formatted = f"{amount:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        return f"€ {formatted}"
    except Exception:
        return "€ 0,00"

def parse_int(value: Any, default: int = 1) -> int:
    """Converte un valore in intero strettamente positivo."""
    if value is None or str(value).strip() == "":
        return default
    try:
        val = int(str(value).strip())
        if val <= 0:
            raise ValueError("La quantità deve essere maggiore di zero.")
        return val
    except ValueError:
        raise ValueError("La quantità deve essere un intero maggiore di zero.")


def parse_date(value: Any, allow_empty: bool = True) -> str:
    """Valida e normalizza una data convertendola sempre nel formato ISO YYYY-MM-DD."""
    text = str(value or "").strip()
    if not text and allow_empty:
        return ""
    
    formats = [
        "%Y-%m-%d",
        "%d/%m/%Y",
        "%d-%m-%Y",
        "%Y/%m/%d",
        "%d.%m.%Y"
    ]
    
    for fmt in formats:
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
            
    raise ValueError(f"Formato data non supportato: {text}")
            
def format_month_label(month_key: Any) -> str:
    """Formatta una stringa mese ISO (YYYY-MM) o data nel formato abbreviato compatto (es. 'SET \\'26')."""
    text = str(month_key or "").strip()
    if not text:
        return ""
    if len(text) >= 7 and "-" in text:
        year_part = text[:4]
        month_part = text[5:7]
        short_name = MONTHS_SHORT_IT.get(month_part, month_part)
        return f"{short_name} '{year_part[2:]}"
    return text

def calculate_profit(purchased_price: Any, sold_price: Optional[Any], fees: Any = Decimal("0"), quantity: int = 1) -> Tuple[Decimal, Decimal]:
    """
    Calcola il profitto totale netto e il ROI % per un articolo venduto.

    - purchased_price : prezzo di acquisto UNITARIO
    - sold_price      : prezzo di vendita TOTALE incassato (non per unità)
    - fees            : commissioni / spese di spedizione TOTALI della vendita (non per unità)
    - quantity        : quantità venduta

    Profitto totale = sold_price - (purchased_price * quantity) - fees
    ROI % = (Profitto totale / (purchased_price * quantity)) * 100
    """
    purchased = purchased_price if isinstance(purchased_price, Decimal) else Decimal(str(purchased_price))
    sold = sold_price if isinstance(sold_price, Decimal) else (Decimal(str(sold_price)) if sold_price is not None else None)
    fee_amount = fees if isinstance(fees, Decimal) else Decimal(str(fees))
    qty = Decimal(str(quantity)) if quantity and quantity > 0 else Decimal("1")

    if sold is None or purchased < 0:
        return Decimal("0"), Decimal("0")

    total_cost = purchased * qty
    total_profit = sold - total_cost - fee_amount
    roi = (total_profit / total_cost) * 100 if total_cost > 0 else Decimal("0")
    return total_profit.quantize(Decimal("0.01")), roi.quantize(Decimal("0.01"))

def normalize_header(value: Any) -> str:
    """Normalizza le intestazioni di tabelle o fogli Excel."""
    text = str(value or "").strip().upper().replace("_", " ")
    return re.sub(r"\s+", " ", text)


def read_spreadsheet_raw(filepath: str) -> str:
    """
    Legge qualsiasi file .xlsx o .csv e restituisce il contenuto come testo grezzo
    tabulare (tab-separated) da passare all'AI per l'estrazione dei prodotti.
    Funziona con qualsiasi struttura di colonne, lingue e formati.
    """
    path = str(filepath)
    lines: List[str] = []

    if path.lower().endswith(".csv"):
        # Lettura CSV con rilevamento automatico del separatore
        with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
            sample = f.read(4096)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        with open(path, newline="", encoding="utf-8-sig", errors="replace") as f:
            reader = csv.reader(f, dialect)
            for row in reader:
                cleaned = [str(c).strip() for c in row]
                while cleaned and not cleaned[-1]:
                    cleaned.pop()
                if any(cleaned):
                    lines.append("\t".join(cleaned))
    else:
        # Lettura XLSX/XLSM con openpyxl
        from openpyxl import load_workbook
        wb = load_workbook(path, data_only=True, read_only=True)
        for sheet in wb.worksheets:
            sheet_title = sheet.title.strip()
            if sheet_title:
                lines.append(f"=== Foglio: {sheet_title} ===")
            for row in sheet.iter_rows(values_only=True):
                cleaned = [str(c).strip() if c is not None else "" for c in row]
                while cleaned and not cleaned[-1]:
                    cleaned.pop()
                if any(c for c in cleaned):
                    lines.append("\t".join(cleaned))
        wb.close()

    return "\n".join(lines)


def export_to_csv(filepath: str, products: List[dict]) -> int:
    """Esporta una lista di prodotti in un file CSV formattato per Excel."""
    headers = [
        "ID", "NOME", "TAGLIA", "QUANTITÀ", "PREZZO ACQUISTO", "LUOGO ACQUISTO",
        "LUOGO VENDITA", "PREZZO VENDITA TOTALE (PREVISTO)", "PREZZO VENDITA TOTALE (EFFETTIVO)",
        "COMMISSIONI/SPESE", "DATA VENDITA", "STATO", "DATA USCITA", "NOTE"
    ]
    with open(filepath, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(headers)
        for row in products:
            p_price = row.get("purchased_price", 0.0)
            p_str = f"{p_price:.2f}".replace(".", ",") if p_price is not None else ""
            s_price = row.get("selling_price")
            s_str = f"{s_price:.2f}".replace(".", ",") if s_price is not None else ""
            sold_price = row.get("sold_price")
            sold_str = f"{sold_price:.2f}".replace(".", ",") if sold_price is not None else ""
            fees = row.get("fees", 0.0)
            fees_str = f"{fees:.2f}".replace(".", ",") if fees is not None else "0,00"

            writer.writerow([
                row.get("id", ""),
                row.get("name", ""),
                row.get("size", ""),
                row.get("quantity", 1),
                p_str,
                row.get("purchased_place", ""),
                row.get("selling_place", ""),
                s_str,
                sold_str,
                fees_str,
                row.get("sold_date", ""),
                row.get("status", ""),
                row.get("release_date", ""),
                row.get("notes", "")
            ])
    return len(products)

def apply_window_icon(window: Any) -> None:
    """Imposta l'icona personalizzata ad alta definizione su qualsiasi finestra (CTk o CTkToplevel)."""
    try:
        import sys
        import ctypes
        from pathlib import Path
        from PIL import Image, ImageTk

        ico_path = Path(getattr(sys, "_MEIPASS", PROJECT_DIR)) / "assets" / "app_icon.ico"
        if not ico_path.exists():
            ico_path = PROJECT_DIR / "assets" / "app_icon.ico"
        if not ico_path.exists():
            ico_path = PROJECT_DIR / "app_icon.ico"

        logo_path = Path(getattr(sys, "_MEIPASS", PROJECT_DIR)) / "assets" / "app_logo.png"
        if not logo_path.exists():
            logo_path = PROJECT_DIR / "assets" / "app_logo.png"

        # 1. Fallback Tkinter standard iconbitmap
        if ico_path.exists():
            try:
                window.iconbitmap(str(ico_path))
            except Exception:
                pass

        # 2. Registra le icone multi-risoluzione native (16, 24, 32, 48, 64, 128, 256) per Tkinter
        source_img_path = logo_path if logo_path.exists() else ico_path
        if source_img_path.exists():
            try:
                base_img = Image.open(source_img_path).convert("RGBA")
                photos = []
                for sz in (16, 24, 32, 48, 64, 128, 256):
                    resized = base_img.resize((sz, sz), Image.Resampling.LANCZOS)
                    photos.append(ImageTk.PhotoImage(resized))
                window._icon_photos_keepalive = photos
                window.iconphoto(False, *photos)
            except Exception:
                pass

        # 3. Chiamata Win32 diretta WM_SETICON per rendere la barra delle applicazioni perfettamente nitida
        if ico_path.exists() and sys.platform == "win32":
            try:
                window.update_idletasks()
                hwnd = ctypes.windll.user32.GetAncestor(window.winfo_id(), 2)  # GA_ROOT = 2
                h_icon_big = ctypes.windll.user32.LoadImageW(None, str(ico_path), 1, 0, 0, 0x0010 | 0x0040)
                h_icon_small = ctypes.windll.user32.LoadImageW(None, str(ico_path), 1, 16, 16, 0x0010)
                WM_SETICON = 0x0080
                if h_icon_big:
                    ctypes.windll.user32.SendMessageW(hwnd, WM_SETICON, 1, h_icon_big)
                if h_icon_small:
                    ctypes.windll.user32.SendMessageW(hwnd, WM_SETICON, 0, h_icon_small)
            except Exception:
                pass
    except Exception:
        pass


def apply_dark_titlebar(window) -> None:
    """Forza la titlebar di Windows 10/11 in modalità scura per matchare il tema."""
    try:
        import ctypes
        window.update_idletasks()
        hwnd = ctypes.windll.user32.GetAncestor(window.winfo_id(), 2) # GA_ROOT = 2
        value = ctypes.c_int(1)
        # Try Windows 11
        ctypes.windll.dwmapi.DwmSetWindowAttribute(hwnd, 20, ctypes.byref(value), ctypes.sizeof(value))
        # Try Windows 10
        ctypes.windll.dwmapi.DwmSetWindowAttribute(hwnd, 19, ctypes.byref(value), ctypes.sizeof(value))
    except Exception:
        pass


def center_window(window: Any, width: int, height: int) -> None:
    """Centra una finestra (CTkToplevel o CTk) perfettamente allo schermo su qualsiasi monitor Windows, gestendo correttamente il DPI scaling."""
    try:
        import ctypes
        screen_w = ctypes.windll.user32.GetSystemMetrics(0)
        screen_h = ctypes.windll.user32.GetSystemMetrics(1)
    except Exception:
        screen_w = window.winfo_screenwidth()
        screen_h = window.winfo_screenheight()

    scaling = getattr(window, "_get_window_scaling", lambda: 1.0)()
    w_physical = round(width * scaling)
    h_physical = round(height * scaling)

    x = max(0, (screen_w - w_physical) // 2)
    y = max(0, (screen_h - h_physical) // 2)

    window.geometry(f"{width}x{height}+{x}+{y}")
    apply_window_icon(window)
    apply_dark_titlebar(window)

