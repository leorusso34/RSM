#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Russo Stock Manager

Gestionale desktop per inventario e pre-ordini.
Dati persistenti su SQLite, creato automaticamente al primo avvio.
GUI realizzata con Tkinter/ttk, inclusi nella distribuzione standard di Python.
"""

from __future__ import annotations

import csv
import re
import shutil
import sqlite3
import sys
import tkinter as tk
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from tkinter import filedialog, messagebox, ttk
from typing import Any, Optional

APP_NAME = "Russo Stock Manager"
APP_VERSION = "1.0.0"
STATUS_STOCK = "In Stock"
STATUS_PREORDER = "Pre-order"


def app_directory() -> Path:
    """Restituisce la cartella del file .py o dell'eseguibile PyInstaller."""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


BASE_DIR = app_directory()
DB_PATH = BASE_DIR / "russo_stock.db"


def parse_money(value: Any, allow_empty: bool = False) -> Optional[float]:
    """Converte valori monetari italiani/internazionali in float."""
    if value is None or str(value).strip() == "":
        if allow_empty:
            return None
        raise ValueError("Il prezzo di acquisto è obbligatorio.")

    text = str(value).strip()
    text = re.sub(r"[^0-9,\.\-]", "", text)
    if not text or text in {"-", ".", ","}:
        if allow_empty:
            return None
        raise ValueError("Inserisci un prezzo valido.")

    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    else:
        text = text.replace(",", ".")

    try:
        amount = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError("Inserisci un prezzo valido.") from exc
    if amount < 0:
        raise ValueError("Il prezzo non può essere negativo.")
    return float(amount.quantize(Decimal("0.01")))


def money(value: Any) -> str:
    amount = Decimal(str(value or 0)).quantize(Decimal("0.01"))
    formatted = f"{amount:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    return f"€ {formatted}"


class InventoryDatabase:
    def __init__(self, path: Path = DB_PATH) -> None:
        self.path = path
        self.connection = sqlite3.connect(self.path)
        self.connection.row_factory = sqlite3.Row
        self.connection.execute("PRAGMA foreign_keys = ON")
        self.connection.execute("PRAGMA journal_mode = WAL")
        self._create_schema()

    def _create_schema(self) -> None:
        self.connection.executescript(
            f"""
            CREATE TABLE IF NOT EXISTS products (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                name TEXT NOT NULL COLLATE NOCASE,
                size TEXT NOT NULL DEFAULT '',
                purchased_price REAL NOT NULL CHECK (purchased_price >= 0),
                purchased_place TEXT NOT NULL DEFAULT '',
                selling_place TEXT NOT NULL DEFAULT '',
                selling_price REAL CHECK (selling_price IS NULL OR selling_price >= 0),
                status TEXT NOT NULL DEFAULT '{STATUS_STOCK}'
                    CHECK (status IN ('{STATUS_STOCK}', '{STATUS_PREORDER}')),
                release_date TEXT NOT NULL DEFAULT '',
                notes TEXT NOT NULL DEFAULT '',
                created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
                updated_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
            );
            CREATE INDEX IF NOT EXISTS idx_products_status ON products(status);
            CREATE INDEX IF NOT EXISTS idx_products_name ON products(name);
            """
        )
        self.connection.commit()

    def list_products(self, status: str, search: str = "") -> list[sqlite3.Row]:
        needle = f"%{search.strip()}%"
        return self.connection.execute(
            """
            SELECT * FROM products
            WHERE status = ? AND (
                name LIKE ? OR size LIKE ? OR purchased_place LIKE ? OR
                selling_place LIKE ? OR release_date LIKE ? OR notes LIKE ?
            )
            ORDER BY name COLLATE NOCASE, id DESC
            """,
            (status, needle, needle, needle, needle, needle, needle),
        ).fetchall()

    def get(self, product_id: int) -> Optional[sqlite3.Row]:
        return self.connection.execute(
            "SELECT * FROM products WHERE id = ?", (product_id,)
        ).fetchone()

    def save(self, data: dict[str, Any], product_id: Optional[int] = None) -> int:
        fields = (
            "name", "size", "purchased_price", "purchased_place",
            "selling_place", "selling_price", "status", "release_date", "notes"
        )
        values = tuple(data[field] for field in fields)
        with self.connection:
            if product_id is None:
                placeholders = ", ".join("?" for _ in fields)
                cursor = self.connection.execute(
                    f"INSERT INTO products ({', '.join(fields)}) VALUES ({placeholders})",
                    values,
                )
                return int(cursor.lastrowid)
            assignments = ", ".join(f"{field} = ?" for field in fields)
            self.connection.execute(
                f"UPDATE products SET {assignments}, updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                values + (product_id,),
            )
            return product_id

    def delete(self, product_id: int) -> None:
        with self.connection:
            self.connection.execute("DELETE FROM products WHERE id = ?", (product_id,))

    def move_to_stock(self, product_id: int) -> None:
        with self.connection:
            self.connection.execute(
                "UPDATE products SET status = ?, release_date = '', updated_at = CURRENT_TIMESTAMP WHERE id = ?",
                (STATUS_STOCK, product_id),
            )

    def totals(self) -> dict[str, Any]:
        row = self.connection.execute(
            """
            SELECT
                COALESCE(SUM(CASE WHEN status = ? THEN purchased_price ELSE 0 END), 0) AS stock_total,
                SUM(CASE WHEN status = ? THEN 1 ELSE 0 END) AS stock_count,
                SUM(CASE WHEN status = ? THEN 1 ELSE 0 END) AS preorder_count,
                COUNT(*) AS all_count,
                COALESCE(SUM(purchased_price), 0) AS committed_total
            FROM products
            """,
            (STATUS_STOCK, STATUS_STOCK, STATUS_PREORDER),
        ).fetchone()
        return dict(row)

    def insert_many(self, records: list[dict[str, Any]]) -> int:
        count = 0
        with self.connection:
            for record in records:
                self.save(record)
                count += 1
        return count

    def close(self) -> None:
        self.connection.close()


class ProductDialog(tk.Toplevel):
    def __init__(
        self,
        parent: tk.Misc,
        initial_status: str,
        record: Optional[sqlite3.Row] = None,
    ) -> None:
        super().__init__(parent)
        self.result: Optional[dict[str, Any]] = None
        self.record = record
        self.title("Modifica prodotto" if record else "Nuovo prodotto")
        self.geometry("600x680")
        self.minsize(560, 620)
        self.configure(bg="#111827")
        self.transient(parent)
        self.grab_set()

        self.vars = {
            "name": tk.StringVar(value=record["name"] if record else ""),
            "size": tk.StringVar(value=record["size"] if record else ""),
            "purchased_price": tk.StringVar(value=f'{record["purchased_price"]:.2f}' if record else ""),
            "purchased_place": tk.StringVar(value=record["purchased_place"] if record else ""),
            "selling_place": tk.StringVar(value=record["selling_place"] if record else ""),
            "selling_price": tk.StringVar(
                value=f'{record["selling_price"]:.2f}' if record and record["selling_price"] is not None else ""
            ),
            "status": tk.StringVar(value=record["status"] if record else initial_status),
            "release_date": tk.StringVar(value=record["release_date"] if record else ""),
        }
        self._build()
        self.protocol("WM_DELETE_WINDOW", self.destroy)
        self.bind("<Escape>", lambda _event: self.destroy())
        self.bind("<Control-Return>", lambda _event: self._submit())
        self.after(80, self._focus_name)

    def _focus_name(self) -> None:
        self.name_entry.focus_set()

    def _build(self) -> None:
        outer = tk.Frame(self, bg="#111827", padx=26, pady=22)
        outer.pack(fill="both", expand=True)

        tk.Label(
            outer, text="Dettagli prodotto", bg="#111827", fg="#f8fafc",
            font=("Segoe UI", 20, "bold"), anchor="w"
        ).pack(fill="x")
        tk.Label(
            outer, text="I campi con * sono obbligatori", bg="#111827", fg="#94a3b8",
            font=("Segoe UI", 10), anchor="w"
        ).pack(fill="x", pady=(2, 16))

        form = tk.Frame(outer, bg="#111827")
        form.pack(fill="both", expand=True)
        form.columnconfigure(0, weight=1)
        form.columnconfigure(1, weight=1)

        self.name_entry = self._field(form, "NAME *", "name", 0, 0, 2)
        self._field(form, "SIZE", "size", 1, 0)
        self._field(form, "PURCHASED PRICE *", "purchased_price", 1, 1)
        self._field(form, "PURCHASED PLACE", "purchased_place", 2, 0)
        self._field(form, "SELLING PLACE", "selling_place", 2, 1)
        self._field(form, "SELLING PRICE", "selling_price", 3, 0)

        label = tk.Label(form, text="STATO *", bg="#111827", fg="#cbd5e1", font=("Segoe UI", 9, "bold"))
        label.grid(row=6, column=1, sticky="w", padx=(8, 0), pady=(7, 4))
        status_box = ttk.Combobox(
            form, textvariable=self.vars["status"], values=(STATUS_STOCK, STATUS_PREORDER),
            state="readonly", font=("Segoe UI", 10)
        )
        status_box.grid(row=7, column=1, sticky="ew", padx=(8, 0), ipady=5)
        status_box.bind("<<ComboboxSelected>>", lambda _e: self._toggle_release())

        self.release_label = tk.Label(
            form, text="DATA USCITA (AAAA-MM-GG)", bg="#111827", fg="#cbd5e1",
            font=("Segoe UI", 9, "bold")
        )
        self.release_label.grid(row=8, column=0, sticky="w", pady=(13, 4))
        self.release_entry = ttk.Entry(form, textvariable=self.vars["release_date"], font=("Segoe UI", 10))
        self.release_entry.grid(row=9, column=0, sticky="ew", padx=(0, 8), ipady=5)

        tk.Label(
            form, text="NOTE", bg="#111827", fg="#cbd5e1", font=("Segoe UI", 9, "bold")
        ).grid(row=10, column=0, columnspan=2, sticky="w", pady=(13, 4))
        self.notes = tk.Text(
            form, height=5, wrap="word", bg="#1f2937", fg="#f8fafc", insertbackground="#f8fafc",
            relief="flat", padx=10, pady=8, font=("Segoe UI", 10)
        )
        self.notes.grid(row=11, column=0, columnspan=2, sticky="nsew")
        if self.record:
            self.notes.insert("1.0", self.record["notes"])
        form.rowconfigure(11, weight=1)

        buttons = tk.Frame(outer, bg="#111827")
        buttons.pack(fill="x", pady=(20, 0))
        ttk.Button(buttons, text="Annulla", style="Secondary.TButton", command=self.destroy).pack(side="right")
        ttk.Button(buttons, text="Salva prodotto", style="Accent.TButton", command=self._submit).pack(side="right", padx=(0, 10))
        self._toggle_release()

    def _field(self, parent: tk.Widget, label: str, key: str, group_row: int, column: int, span: int = 1) -> ttk.Entry:
        label_row = group_row * 2
        entry_row = label_row + 1
        padx = (0, 8) if column == 0 and span == 1 else ((8, 0) if column == 1 else 0)
        tk.Label(parent, text=label, bg="#111827", fg="#cbd5e1", font=("Segoe UI", 9, "bold")).grid(
            row=label_row, column=column, columnspan=span, sticky="w", padx=padx, pady=(7, 4)
        )
        entry = ttk.Entry(parent, textvariable=self.vars[key], font=("Segoe UI", 10))
        entry.grid(row=entry_row, column=column, columnspan=span, sticky="ew", padx=padx, ipady=5)
        return entry

    def _toggle_release(self) -> None:
        state = "normal" if self.vars["status"].get() == STATUS_PREORDER else "disabled"
        self.release_entry.configure(state=state)
        self.release_label.configure(fg="#cbd5e1" if state == "normal" else "#64748b")
        if state == "disabled":
            self.vars["release_date"].set("")

    def _submit(self) -> None:
        try:
            name = self.vars["name"].get().strip()
            if not name:
                raise ValueError("Il nome del prodotto è obbligatorio.")
            purchase = parse_money(self.vars["purchased_price"].get())
            sale = parse_money(self.vars["selling_price"].get(), allow_empty=True)
            release = self.vars["release_date"].get().strip()
            if release:
                datetime.strptime(release, "%Y-%m-%d")
        except ValueError as exc:
            message = "La data di uscita deve avere formato AAAA-MM-GG." if "time data" in str(exc) else str(exc)
            messagebox.showerror("Dati non validi", message, parent=self)
            return

        self.result = {
            "name": name,
            "size": self.vars["size"].get().strip(),
            "purchased_price": purchase,
            "purchased_place": self.vars["purchased_place"].get().strip(),
            "selling_place": self.vars["selling_place"].get().strip(),
            "selling_price": sale,
            "status": self.vars["status"].get(),
            "release_date": release,
            "notes": self.notes.get("1.0", "end").strip(),
        }
        self.destroy()


class InventoryApp(tk.Tk):
    COLUMNS = (
        "id", "name", "size", "purchased_price", "purchased_place",
        "selling_place", "selling_price", "release_date"
    )
    HEADINGS = {
        "id": "ID", "name": "NAME", "size": "SIZE", "purchased_price": "PURCHASED PRICE",
        "purchased_place": "PURCHASED PLACE", "selling_place": "SELLING PLACE",
        "selling_price": "SELLING PRICE", "release_date": "DATA USCITA"
    }

    def __init__(self) -> None:
        super().__init__()
        self.db = InventoryDatabase()
        self.title(f"{APP_NAME} {APP_VERSION}")
        self.geometry("1380x820")
        self.minsize(1040, 660)
        self.configure(bg="#0b1120")
        self.search_var = tk.StringVar()
        self.status_var = tk.StringVar(value="Pronto")
        self.current_sort: dict[str, tuple[str, bool]] = {}
        self._setup_style()
        self._build_ui()
        self.refresh()
        self.protocol("WM_DELETE_WINDOW", self._on_close)
        self.bind("<Control-n>", lambda _e: self.add_product())
        self.bind("<Control-f>", lambda _e: self.search_entry.focus_set())
        self.bind("<F5>", lambda _e: self.refresh())
        self.after(150, self.search_entry.focus_set)

    def _setup_style(self) -> None:
        style = ttk.Style(self)
        try:
            style.theme_use("clam")
        except tk.TclError:
            pass
        style.configure("TFrame", background="#0b1120")
        style.configure("Card.TFrame", background="#111827")
        style.configure("TLabel", background="#0b1120", foreground="#e5e7eb", font=("Segoe UI", 10))
        style.configure("Muted.TLabel", foreground="#94a3b8")
        style.configure("Title.TLabel", foreground="#f8fafc", font=("Segoe UI", 24, "bold"))
        style.configure("KpiValue.TLabel", background="#111827", foreground="#f8fafc", font=("Segoe UI", 21, "bold"))
        style.configure("KpiLabel.TLabel", background="#111827", foreground="#94a3b8", font=("Segoe UI", 9, "bold"))
        style.configure("TButton", background="#1f2937", foreground="#e5e7eb", borderwidth=0, padding=(14, 9), font=("Segoe UI", 9, "bold"))
        style.map("TButton", background=[("active", "#374151")])
        style.configure("Accent.TButton", background="#2563eb", foreground="white")
        style.map("Accent.TButton", background=[("active", "#1d4ed8")])
        style.configure("Danger.TButton", background="#991b1b", foreground="white")
        style.map("Danger.TButton", background=[("active", "#b91c1c")])
        style.configure("Secondary.TButton", background="#374151", foreground="#f8fafc")
        style.configure("TEntry", fieldbackground="#1f2937", foreground="#f8fafc", insertcolor="#f8fafc", borderwidth=0, padding=8)
        style.configure("TCombobox", fieldbackground="#1f2937", background="#1f2937", foreground="#f8fafc", arrowcolor="#cbd5e1", padding=6)
        style.map("TCombobox", fieldbackground=[("readonly", "#1f2937")], foreground=[("readonly", "#f8fafc")])
        style.configure("TNotebook", background="#0b1120", borderwidth=0)
        style.configure("TNotebook.Tab", background="#111827", foreground="#94a3b8", padding=(20, 11), font=("Segoe UI", 10, "bold"))
        style.map("TNotebook.Tab", background=[("selected", "#2563eb")], foreground=[("selected", "white")])
        style.configure("Treeview", background="#111827", fieldbackground="#111827", foreground="#e5e7eb", rowheight=34, borderwidth=0, font=("Segoe UI", 9))
        style.map("Treeview", background=[("selected", "#1d4ed8")], foreground=[("selected", "white")])
        style.configure("Treeview.Heading", background="#1f2937", foreground="#cbd5e1", relief="flat", padding=9, font=("Segoe UI", 9, "bold"))
        style.map("Treeview.Heading", background=[("active", "#334155")])

    def _build_ui(self) -> None:
        container = ttk.Frame(self, padding=24)
        container.pack(fill="both", expand=True)

        header = ttk.Frame(container)
        header.pack(fill="x", pady=(0, 18))
        title_box = ttk.Frame(header)
        title_box.pack(side="left")
        ttk.Label(title_box, text="RUSSO STOCK", style="Title.TLabel").pack(anchor="w")
        ttk.Label(title_box, text="Inventario, pre-ordini e valore del magazzino", style="Muted.TLabel").pack(anchor="w")

        actions = ttk.Frame(header)
        actions.pack(side="right", pady=4)
        ttk.Button(actions, text="Importa Excel", command=self.import_excel).pack(side="left", padx=4)
        ttk.Button(actions, text="Esporta CSV", command=self.export_csv).pack(side="left", padx=4)
        ttk.Button(actions, text="Backup DB", command=self.backup_database).pack(side="left", padx=4)
        ttk.Button(actions, text="+ Nuovo prodotto", style="Accent.TButton", command=self.add_product).pack(side="left", padx=(10, 0))

        kpis = ttk.Frame(container)
        kpis.pack(fill="x", pady=(0, 18))
        self.kpi_vars: dict[str, tk.StringVar] = {}
        cards = [
            ("stock_total", "STOCK TOTAL", "#2563eb"),
            ("stock_count", "PEZZI IN MAGAZZINO", "#059669"),
            ("preorder_count", "PRE-ORDINI", "#d97706"),
            ("all_count", "PEZZI TOTALI", "#7c3aed"),
        ]
        for index, (key, label, color) in enumerate(cards):
            kpis.columnconfigure(index, weight=1)
            card = tk.Frame(kpis, bg="#111827", highlightbackground=color, highlightthickness=1, padx=18, pady=14)
            card.grid(row=0, column=index, sticky="nsew", padx=(0 if index == 0 else 7, 0 if index == len(cards) - 1 else 7))
            var = tk.StringVar(value="-")
            self.kpi_vars[key] = var
            tk.Label(card, text=label, bg="#111827", fg="#94a3b8", font=("Segoe UI", 9, "bold")).pack(anchor="w")
            tk.Label(card, textvariable=var, bg="#111827", fg="#f8fafc", font=("Segoe UI", 21, "bold")).pack(anchor="w", pady=(6, 0))

        toolbar = ttk.Frame(container)
        toolbar.pack(fill="x", pady=(0, 12))
        ttk.Label(toolbar, text="Cerca:").pack(side="left")
        self.search_entry = ttk.Entry(toolbar, textvariable=self.search_var, width=38)
        self.search_entry.pack(side="left", padx=(8, 6), ipady=2)
        self.search_var.trace_add("write", lambda *_args: self.refresh_tables())
        ttk.Button(toolbar, text="Azzera", command=lambda: self.search_var.set("")).pack(side="left")
        ttk.Label(toolbar, text="Doppio clic su una riga per modificarla", style="Muted.TLabel").pack(side="right")

        self.notebook = ttk.Notebook(container)
        self.notebook.pack(fill="both", expand=True)
        self.trees: dict[str, ttk.Treeview] = {}
        self._create_tab(STATUS_STOCK, "In Magazzino")
        self._create_tab(STATUS_PREORDER, "Pre-ordini")
        self.notebook.bind("<<NotebookTabChanged>>", lambda _e: self._update_action_states())

        bottom = ttk.Frame(container)
        bottom.pack(fill="x", pady=(12, 0))
        self.edit_button = ttk.Button(bottom, text="Modifica", command=self.edit_selected)
        self.edit_button.pack(side="left", padx=(0, 6))
        self.delete_button = ttk.Button(bottom, text="Elimina", style="Danger.TButton", command=self.delete_selected)
        self.delete_button.pack(side="left", padx=6)
        self.move_button = ttk.Button(bottom, text="Converti in In Magazzino", style="Accent.TButton", command=self.move_selected_to_stock)
        self.move_button.pack(side="left", padx=6)
        ttk.Label(bottom, textvariable=self.status_var, style="Muted.TLabel").pack(side="right")

    def _create_tab(self, status: str, title: str) -> None:
        frame = ttk.Frame(self.notebook, padding=(0, 12, 0, 0))
        self.notebook.add(frame, text=title)
        tree_frame = ttk.Frame(frame, style="Card.TFrame")
        tree_frame.pack(fill="both", expand=True)
        tree = ttk.Treeview(tree_frame, columns=self.COLUMNS, show="headings", selectmode="browse")
        widths = {"id": 55, "name": 300, "size": 90, "purchased_price": 135, "purchased_place": 160, "selling_place": 145, "selling_price": 125, "release_date": 115}
        for column in self.COLUMNS:
            tree.heading(column, text=self.HEADINGS[column], command=lambda c=column, s=status: self.sort_tree(s, c))
            tree.column(column, width=widths[column], minwidth=50, anchor="e" if "price" in column else "w", stretch=column == "name")
        yscroll = ttk.Scrollbar(tree_frame, orient="vertical", command=tree.yview)
        xscroll = ttk.Scrollbar(tree_frame, orient="horizontal", command=tree.xview)
        tree.configure(yscrollcommand=yscroll.set, xscrollcommand=xscroll.set)
        tree.grid(row=0, column=0, sticky="nsew")
        yscroll.grid(row=0, column=1, sticky="ns")
        xscroll.grid(row=1, column=0, sticky="ew")
        tree_frame.rowconfigure(0, weight=1)
        tree_frame.columnconfigure(0, weight=1)
        tree.bind("<Double-1>", lambda _e: self.edit_selected())
        tree.bind("<Delete>", lambda _e: self.delete_selected())
        tree.bind("<<TreeviewSelect>>", lambda _e: self._update_action_states())
        self.trees[status] = tree

    def current_status(self) -> str:
        return STATUS_STOCK if self.notebook.index("current") == 0 else STATUS_PREORDER

    def current_tree(self) -> ttk.Treeview:
        return self.trees[self.current_status()]

    def selected_id(self) -> Optional[int]:
        selection = self.current_tree().selection()
        return int(selection[0]) if selection else None

    def refresh(self) -> None:
        self.refresh_tables()
        self.refresh_totals()
        self._update_action_states()

    def refresh_tables(self) -> None:
        query = self.search_var.get()
        for status, tree in self.trees.items():
            selected = tree.selection()
            selected_id = selected[0] if selected else None
            tree.delete(*tree.get_children())
            for row in self.db.list_products(status, query):
                values = (
                    row["id"], row["name"], row["size"], money(row["purchased_price"]),
                    row["purchased_place"], row["selling_place"],
                    money(row["selling_price"]) if row["selling_price"] is not None else "",
                    row["release_date"],
                )
                tree.insert("", "end", iid=str(row["id"]), values=values)
            if selected_id and tree.exists(selected_id):
                tree.selection_set(selected_id)
        self.status_var.set("Vista aggiornata")

    def refresh_totals(self) -> None:
        totals = self.db.totals()
        self.kpi_vars["stock_total"].set(money(totals["stock_total"]))
        self.kpi_vars["stock_count"].set(str(totals["stock_count"]))
        self.kpi_vars["preorder_count"].set(str(totals["preorder_count"]))
        self.kpi_vars["all_count"].set(str(totals["all_count"]))

    def _update_action_states(self) -> None:
        has_selection = self.selected_id() is not None
        state = "normal" if has_selection else "disabled"
        self.edit_button.configure(state=state)
        self.delete_button.configure(state=state)
        self.move_button.configure(state=state if has_selection and self.current_status() == STATUS_PREORDER else "disabled")

    def add_product(self) -> None:
        dialog = ProductDialog(self, self.current_status())
        self.wait_window(dialog)
        if dialog.result:
            new_id = self.db.save(dialog.result)
            target = dialog.result["status"]
            self.refresh()
            if target != self.current_status():
                self.notebook.select(0 if target == STATUS_STOCK else 1)
            tree = self.trees[target]
            if tree.exists(str(new_id)):
                tree.selection_set(str(new_id))
                tree.see(str(new_id))
            self.status_var.set("Prodotto aggiunto")

    def edit_selected(self) -> None:
        product_id = self.selected_id()
        if product_id is None:
            return
        record = self.db.get(product_id)
        if record is None:
            self.refresh()
            return
        dialog = ProductDialog(self, record["status"], record)
        self.wait_window(dialog)
        if dialog.result:
            self.db.save(dialog.result, product_id)
            target = dialog.result["status"]
            self.refresh()
            self.notebook.select(0 if target == STATUS_STOCK else 1)
            tree = self.trees[target]
            if tree.exists(str(product_id)):
                tree.selection_set(str(product_id))
                tree.see(str(product_id))
            self.status_var.set("Prodotto modificato")

    def delete_selected(self) -> None:
        product_id = self.selected_id()
        if product_id is None:
            return
        record = self.db.get(product_id)
        if record and messagebox.askyesno(
            "Conferma eliminazione",
            f'Eliminare definitivamente "{record["name"]}"?\n\nQuesta operazione non può essere annullata.',
            parent=self,
        ):
            self.db.delete(product_id)
            self.refresh()
            self.status_var.set("Prodotto eliminato")

    def move_selected_to_stock(self) -> None:
        product_id = self.selected_id()
        if product_id is None or self.current_status() != STATUS_PREORDER:
            return
        record = self.db.get(product_id)
        if record and messagebox.askyesno(
            "Converti pre-ordine",
            f'Spostare "{record["name"]}" in In Magazzino?',
            parent=self,
        ):
            self.db.move_to_stock(product_id)
            self.refresh()
            self.notebook.select(0)
            tree = self.trees[STATUS_STOCK]
            if tree.exists(str(product_id)):
                tree.selection_set(str(product_id))
                tree.see(str(product_id))
            self.status_var.set("Pre-ordine convertito in prodotto a magazzino")

    def sort_tree(self, status: str, column: str) -> None:
        tree = self.trees[status]
        previous_column, reverse = self.current_sort.get(status, ("", False))
        reverse = not reverse if previous_column == column else False
        self.current_sort[status] = (column, reverse)
        rows = [(tree.set(item, column), item) for item in tree.get_children("")]

        def key(item: tuple[str, str]) -> Any:
            value = item[0]
            if column in {"id", "purchased_price", "selling_price"}:
                try:
                    return parse_money(value, allow_empty=True) or 0
                except ValueError:
                    return 0
            return value.casefold()

        rows.sort(key=key, reverse=reverse)
        for index, (_value, item) in enumerate(rows):
            tree.move(item, "", index)

    def export_csv(self) -> None:
        path = filedialog.asksaveasfilename(
            parent=self, title="Esporta inventario", defaultextension=".csv",
            filetypes=[("File CSV", "*.csv")], initialfile=f"russo_stock_{datetime.now():%Y%m%d}.csv"
        )
        if not path:
            return
        rows = self.db.connection.execute("SELECT * FROM products ORDER BY status, name").fetchall()
        headers = ["ID", "NAME", "SIZE", "PURCHASED PRICE", "PURCHASED PLACE", "SELLING PLACE", "SELLING PRICE", "STATO", "DATA USCITA", "NOTE"]
        with open(path, "w", newline="", encoding="utf-8-sig") as handle:
            writer = csv.writer(handle, delimiter=";")
            writer.writerow(headers)
            for row in rows:
                writer.writerow([
                    row["id"], row["name"], row["size"], f'{row["purchased_price"]:.2f}',
                    row["purchased_place"], row["selling_place"],
                    "" if row["selling_price"] is None else f'{row["selling_price"]:.2f}',
                    row["status"], row["release_date"], row["notes"]
                ])
        self.status_var.set(f"Esportati {len(rows)} prodotti")
        messagebox.showinfo("Esportazione completata", f"File creato:\n{path}", parent=self)

    @staticmethod
    def _normalize_header(value: Any) -> str:
        text = str(value or "").strip().upper().replace("_", " ")
        return re.sub(r"\s+", " ", text)

    def import_excel(self) -> None:
        try:
            from openpyxl import load_workbook
        except ImportError:
            messagebox.showerror(
                "Libreria mancante",
                "Per importare file Excel installa openpyxl:\n\npython -m pip install openpyxl\n\nIl resto del programma funziona senza dipendenze esterne.",
                parent=self,
            )
            return
        path = filedialog.askopenfilename(parent=self, title="Importa da Excel", filetypes=[("Excel", "*.xlsx *.xlsm")])
        if not path:
            return
        try:
            workbook = load_workbook(path, data_only=True, read_only=True)
            aliases = {
                "name": {"NAME", "NOME", "PRODOTTO", "PRODUCT"},
                "size": {"SIZE", "TAGLIA"},
                "purchased_price": {"PURCHASED PRICE", "PREZZO ACQUISTO", "PURCHASE PRICE", "COSTO"},
                "purchased_place": {"PURCHASED PLACE", "LUOGO ACQUISTO", "PURCHASE PLACE", "SITO ACQUISTO"},
                "selling_place": {"SELLING PLACE", "LUOGO VENDITA", "SALE PLACE", "PIATTAFORMA VENDITA"},
                "selling_price": {"SELLING PRICE", "PREZZO VENDITA", "SALE PRICE"},
                "status": {"STATO", "STATUS"},
                "release_date": {"DATA USCITA", "RELEASE DATE", "USCITA"},
                "notes": {"NOTE", "NOTES"},
            }
            records: list[dict[str, Any]] = []
            skipped = 0
            for sheet in workbook.worksheets:
                rows = sheet.iter_rows(values_only=True)
                try:
                    first = next(rows)
                except StopIteration:
                    continue
                normalized = [self._normalize_header(cell) for cell in first]
                mapping: dict[str, int] = {}
                for key, names in aliases.items():
                    for index, header in enumerate(normalized):
                        if header in names:
                            mapping[key] = index
                            break
                if "name" not in mapping or "purchased_price" not in mapping:
                    continue
                default_status = STATUS_PREORDER if "PRE" in sheet.title.upper() else STATUS_STOCK
                for row in rows:
                    def value(key: str, default: Any = "") -> Any:
                        index = mapping.get(key)
                        return row[index] if index is not None and index < len(row) and row[index] is not None else default
                    name = str(value("name")).strip()
                    if not name:
                        continue
                    try:
                        purchase = parse_money(value("purchased_price"))
                        sale = parse_money(value("selling_price"), allow_empty=True)
                    except ValueError:
                        skipped += 1
                        continue
                    raw_status = str(value("status", default_status)).strip().lower()
                    status = STATUS_PREORDER if "pre" in raw_status or "order" in raw_status else STATUS_STOCK
                    release_raw = value("release_date")
                    if isinstance(release_raw, datetime):
                        release = release_raw.strftime("%Y-%m-%d")
                    else:
                        release = str(release_raw or "").strip()
                    records.append({
                        "name": name,
                        "size": str(value("size")).strip(),
                        "purchased_price": purchase,
                        "purchased_place": str(value("purchased_place")).strip(),
                        "selling_place": str(value("selling_place")).strip(),
                        "selling_price": sale,
                        "status": status,
                        "release_date": release if status == STATUS_PREORDER else "",
                        "notes": str(value("notes")).strip(),
                    })
            if not records:
                raise ValueError("Nessuna riga importabile trovata. Verifica che siano presenti almeno le colonne NAME e PURCHASED PRICE.")
            if not messagebox.askyesno(
                "Conferma importazione",
                f"Sono stati trovati {len(records)} prodotti.\nSaranno aggiunti ai dati esistenti. Continuare?",
                parent=self,
            ):
                return
            self.db.insert_many(records)
            self.refresh()
            messagebox.showinfo(
                "Importazione completata",
                f"Prodotti importati: {len(records)}\nRighe ignorate: {skipped}",
                parent=self,
            )
        except Exception as exc:
            messagebox.showerror("Errore di importazione", str(exc), parent=self)

    def backup_database(self) -> None:
        if not DB_PATH.exists():
            return
        target = filedialog.asksaveasfilename(
            parent=self, title="Salva backup database", defaultextension=".db",
            filetypes=[("Database SQLite", "*.db")], initialfile=f"russo_stock_backup_{datetime.now():%Y%m%d_%H%M}.db"
        )
        if target:
            self.db.connection.commit()
            shutil.copy2(DB_PATH, target)
            self.status_var.set("Backup database completato")
            messagebox.showinfo("Backup completato", f"Backup creato:\n{target}", parent=self)

    def _on_close(self) -> None:
        self.db.close()
        self.destroy()


def main() -> None:
    try:
        app = InventoryApp()
        app.mainloop()
    except sqlite3.Error as exc:
        root = tk.Tk()
        root.withdraw()
        messagebox.showerror("Errore database", f"Impossibile avviare {APP_NAME}:\n{exc}")
        root.destroy()
        raise


if __name__ == "__main__":
    main()
