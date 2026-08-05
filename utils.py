#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Utility di formattazione, calcolo e I/O per Russo Stock Manager."""

import csv
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from typing import Any, Optional, Tuple, List, Dict

from config import STATUS_PREORDER, STATUS_SOLD, STATUS_STOCK

def parse_money(value: Any, allow_empty: bool = False) -> Optional[float]:
    """Converte valori monetari italiani/internazionali in float."""
    if value is None or str(value).strip() == "":
        if allow_empty:
            return None
        raise ValueError("Valore monetario obbligatorio.")

    text = str(value).strip()
    text = re.sub(r"[^0-9,\.\-]", "", text)
    if not text or text in {"-", ".", ","}:
        if allow_empty:
            return None
        raise ValueError("Inserisci un importo valido.")

    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    else:
        text = text.replace(",", ".")

    try:
        amount = Decimal(text)
    except InvalidOperation as exc:
        raise ValueError("Inserisci un importo valido.") from exc

    if amount < 0:
        raise ValueError("L'importo non può essere negativo.")

    return float(amount.quantize(Decimal("0.01")))

def money(value: Any) -> str:
    """Formatta un numero come importo monetario in Euro (€ 1.234,56)."""
    if value is None:
        return "€ 0,00"
    try:
        amount = Decimal(str(value)).quantize(Decimal("0.01"))
        formatted = f"{amount:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
        return f"€ {formatted}"
    except Exception:
        return "€ 0,00"

def parse_int(value: Any, default: int = 1) -> int:
    """Converte un valore in intero positivo per la quantità."""
    if value is None or str(value).strip() == "":
        return default
    try:
        val = int(str(value).strip())
        return val if val > 0 else default
    except ValueError:
        return default

def calculate_profit(purchased_price: float, sold_price: Optional[float], fees: float = 0.0, quantity: int = 1) -> Tuple[float, float]:
    """
    Calcola il profitto totale netto e il ROI % per un articolo venduto.
    Profitto unitario = sold_price - purchased_price - fees
    Profitto totale = profitto unitario * quantity
    ROI % = (profitto unitario / purchased_price) * 100
    """
    if sold_price is None or purchased_price <= 0:
        return 0.0, 0.0

    unit_profit = sold_price - purchased_price - fees
    total_profit = unit_profit * quantity
    roi = (unit_profit / purchased_price) * 100 if purchased_price > 0 else 0.0
    return round(total_profit, 2), round(roi, 2)

def normalize_header(value: Any) -> str:
    """Normalizza le intestazioni di tabelle o fogli Excel."""
    text = str(value or "").strip().upper().replace("_", " ")
    return re.sub(r"\s+", " ", text)

def export_to_csv(filepath: str, products: List[dict]) -> int:
    """Esporta una lista di prodotti in un file CSV."""
    headers = [
        "ID", "NOME", "TAGLIA", "QUANTITÀ", "PREZZO ACQUISTO", "LUOGO ACQUISTO",
        "LUOGO VENDITA", "PREZZO VENDITA ESTIMATO", "PREZZO VENDITA EFFETTIVO",
        "COMMISSIONI/SPESE", "DATA VENDITA", "STATO", "DATA USCITA", "NOTE"
    ]
    with open(filepath, "w", newline="", encoding="utf-8-sig") as handle:
        writer = csv.writer(handle, delimiter=";")
        writer.writerow(headers)
        for row in products:
            writer.writerow([
                row.get("id", ""),
                row.get("name", ""),
                row.get("size", ""),
                row.get("quantity", 1),
                f'{row.get("purchased_price", 0.0):.2f}',
                row.get("purchased_place", ""),
                row.get("selling_place", ""),
                "" if row.get("selling_price") is None else f'{row.get("selling_price"):.2f}',
                "" if row.get("sold_price") is None else f'{row.get("sold_price"):.2f}',
                f'{row.get("fees", 0.0):.2f}',
                row.get("sold_date", ""),
                row.get("status", ""),
                row.get("release_date", ""),
                row.get("notes", "")
            ])
    return len(products)

def parse_excel_file(filepath: str) -> Tuple[List[Dict[str, Any]], int]:
    """
    Legge un file Excel e converte le righe in dizionari prodotto.
    Riconosce automaticamente se i prodotti sono sotto una sezione/scritta 'PREORDER'.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(filepath, data_only=True, read_only=True)
    aliases = {
        "name": {"NAME", "NOME", "PRODOTTO", "PRODUCT", "DESCR"},
        "size": {"SIZE", "TAGLIA"},
        "quantity": {"QUANTITY", "QUANTITA", "QTY", "PZ", "PEZZI"},
        "purchased_price": {"PURCHASED PRICE", "PREZZO ACQUISTO", "PURCHASE PRICE", "COSTO", "COST"},
        "purchased_place": {"PURCHASED PLACE", "LUOGO ACQUISTO", "PURCHASE PLACE", "SITO ACQUISTO"},
        "selling_place": {"SELLING PLACE", "LUOGO VENDITA", "SALE PLACE", "PIATTAFORMA VENDITA"},
        "selling_price": {"SELLING PRICE", "PREZZO VENDITA", "SALE PRICE", "PREZZO STIMATO"},
        "sold_price": {"SOLD PRICE", "PREZZO EFFETTIVO", "PREZZO VENDUTO", "INCASSO"},
        "fees": {"FEES", "SPESE", "COMMISSIONI", "COMMISSIONI/SPESE"},
        "sold_date": {"SOLD DATE", "DATA VENDITA", "VENDUTO IL"},
        "status": {"STATO", "STATUS"},
        "release_date": {"DATA USCITA", "RELEASE DATE", "USCITA"},
        "notes": {"NOTE", "NOTES"},
    }

    records: List[Dict[str, Any]] = []
    skipped = 0

    for sheet in workbook.worksheets:
        default_status = STATUS_PREORDER if "PRE" in sheet.title.upper() else STATUS_STOCK
        current_section_status = default_status

        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue

        mapping: Dict[str, int] = {}
        header_found = False

        for row in rows:
            if not row or all(c is None for c in row):
                continue

            # Verifichiamo se l'intera riga contiene una scritta/sezione tipo "PREORDER" o "IN STOCK"
            row_str = " ".join(str(cell or "").upper() for cell in row)
            if "PREORDER" in row_str or "PRE-ORDER" in row_str or "PRE ORDER" in row_str:
                current_section_status = STATUS_PREORDER
            elif "IN STOCK" in row_str or "MAGAZZINO" in row_str or "STOCK" in row_str:
                if "TOTAL" not in row_str and "PRE" not in row_str:
                    current_section_status = STATUS_STOCK

            normalized_row = [normalize_header(cell) for cell in row]

            # Controlla se questa riga è una riga d'intestazione
            temp_mapping: Dict[str, int] = {}
            for key, names in aliases.items():
                for index, header in enumerate(normalized_row):
                    if header in names:
                        temp_mapping[key] = index
                        break

            if "name" in temp_mapping and "purchased_price" in temp_mapping:
                mapping = temp_mapping
                header_found = True
                continue

            if not header_found:
                continue

            # Processa riga dati
            def value(key: str, default: Any = "") -> Any:
                index = mapping.get(key)
                return row[index] if index is not None and index < len(row) and row[index] is not None else default

            name = str(value("name")).strip()
            if not name or name.upper() in {"NAME", "NOME", "PRODOTTO", "TOTAL", "TOTALE"}:
                continue

            try:
                purchase = parse_money(value("purchased_price"))
                sale = parse_money(value("selling_price"), allow_empty=True)
                sold = parse_money(value("sold_price"), allow_empty=True)
                fees = parse_money(value("fees"), allow_empty=True) or 0.0
                qty = parse_int(value("quantity"), default=1)
            except ValueError:
                skipped += 1
                continue

            # Rilevamento dello stato
            raw_status = str(value("status", "")).strip().lower()
            if "vendut" in raw_status or "sold" in raw_status:
                status = STATUS_SOLD
            elif "pre" in raw_status or "order" in raw_status:
                status = STATUS_PREORDER
            elif "stock" in raw_status or "magazzino" in raw_status:
                status = STATUS_STOCK
            else:
                status = current_section_status

            release_raw = value("release_date")
            if isinstance(release_raw, datetime):
                release = release_raw.strftime("%Y-%m-%d")
            else:
                release = str(release_raw or "").strip()

            sold_date_raw = value("sold_date")
            if isinstance(sold_date_raw, datetime):
                sold_date = sold_date_raw.strftime("%Y-%m-%d")
            else:
                sold_date = str(sold_date_raw or "").strip()

            records.append({
                "name": name,
                "size": str(value("size")).strip(),
                "quantity": qty,
                "purchased_price": purchase,
                "purchased_place": str(value("purchased_place")).strip(),
                "selling_place": str(value("selling_place")).strip(),
                "selling_price": sale,
                "sold_price": sold,
                "fees": fees,
                "sold_date": sold_date,
                "status": status,
                "release_date": release if status == STATUS_PREORDER else "",
                "notes": str(value("notes")).strip(),
            })

    return records, skipped
